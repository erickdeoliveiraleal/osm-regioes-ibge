"""Gera regioes_geograficas.xlsx (para converter em Google Sheets).

Abas: Imediatas, Intermediárias e Fontes. Nomes das regiões com link para a relação
no OSM. Municípios e população das imediatas vêm somados do IBGE (a aba de 5.571
municípios deixaria o arquivo grande demais para o upload); os totais das
intermediárias são fórmulas sobre a aba Imediatas.
Uso: python -X utf8 make_sheet.py <pasta_dados>   (mesma pasta usada no compare.py)
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from collections import Counter

d = Path(sys.argv[1])
load = lambda f: json.loads((d / f).read_text(encoding="utf-8"))
rgint = load("rgint.json")
rgi = load("rgi.json")
mun = load("mun.json")
pop_mun = {r["D1C"]: int(r["V"]) for r in load("pop.json")[1:]}
rep = json.load(open("report.json", encoding="utf-8"))
osm_id = {r["codigo"]: r["osm_id"] for r in rep["rows"]}
n_mun = Counter(m["regiao-imediata-id"] for m in mun)
pop_imd = Counter()
for m in mun:
    pop_imd[m["regiao-imediata-id"]] += pop_mun.get(str(m["municipio-id"]), 0)

nome = lambda s: s.replace(" ¿ ", " - ").replace("¿", "-")  # travessão corrompido na API do IBGE
OSM = "https://www.openstreetmap.org/relation/"
FONT = "Arial"
HEAD = Font(name=FONT, bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F4E78")
BODY = Font(name=FONT)
LINK = Font(name=FONT, color="1155CC", underline="single")

wb = Workbook()


def sheet(title, headers, widths, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for c, w in zip(ws[1], widths):
        c.font, c.fill = HEAD, FILL
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = "A2"
    return ws


def link(cell, codigo):
    if codigo in osm_id:
        # Fórmula em vez de hyperlink: o .xlsx fica bem menor para o upload
        texto = str(cell.value).replace('"', '""')
        cell.value = f'=HYPERLINK("{OSM}{osm_id[codigo]}","{texto}")'
        cell.font = LINK


# Imediatas: uma linha por região imediata, como a lista da Wikipédia
ws = sheet("Imediatas", ["Grande região", "UF", "ID intermediária", "Região Geográfica Intermediária",
                         "ID imediata", "Região Geográfica Imediata", "Relação OSM",
                         "Municípios", "População (Censo 2022)"],
           [14, 6, 16, 40, 13, 48, 13, 12, 22], first=True)
for i, r in enumerate(sorted(rgi, key=lambda r: r["id"]), start=2):
    ri = r["regiao-intermediaria"]
    cod, cod_int = str(r["id"]), str(ri["id"])
    ws.append([ri["UF"]["regiao"]["nome"], ri["UF"]["sigla"], ri["id"], nome(ri["nome"]),
               r["id"], nome(r["nome"]), osm_id.get(cod, "não existe no OSM"),
               n_mun[r["id"]], pop_imd[r["id"]]])
    for c in ws[i]:
        c.font = BODY
    link(ws.cell(i, 4), cod_int)
    link(ws.cell(i, 6), cod)
    ws.cell(i, 9).number_format = "#,##0"

# Intermediárias
ws = sheet("Intermediárias", ["Grande região", "UF", "ID intermediária", "Região Geográfica Intermediária",
                              "Relação OSM", "Regiões imediatas", "Municípios", "População (Censo 2022)"],
           [14, 6, 16, 44, 13, 17, 12, 22])
for i, r in enumerate(sorted(rgint, key=lambda r: r["id"]), start=2):
    cod = str(r["id"])
    ws.append([r["UF"]["regiao"]["nome"], r["UF"]["sigla"], r["id"], nome(r["nome"]),
               osm_id.get(cod, "não existe no OSM"),
               f"=COUNTIF(Imediatas!$C:$C,C{i})", f"=SUMIF(Imediatas!$C:$C,C{i},Imediatas!$H:$H)",
               f"=SUMIF(Imediatas!$C:$C,C{i},Imediatas!$I:$I)"])
    for c in ws[i]:
        c.font = BODY
    link(ws.cell(i, 4), cod)
    ws.cell(i, 8).number_format = "#,##0"

ws = sheet("Fontes", ["Dado", "Fonte"], [30, 100])
for row in [
    ("Regiões e municípios", "IBGE, API de Localidades, servicodados.ibge.gov.br/api/v1/localidades"),
    ("População", "IBGE, Censo Demográfico 2022, SIDRA tabela 4709, população residente por município, somada por região imediata"),
    ("Relações OSM", "OpenStreetMap, Overpass API, relações com IBGE:GEOCODIGO, consulta de out. 2026"),
    ("Observação", "Boa Esperança do Norte (MT) foi criado após o Censo 2022; sua população está nos municípios de origem."),
]:
    ws.append(row)
    for c in ws[ws.max_row]:
        c.font = BODY

wb.save("regioes_geograficas.xlsx")
print(Path("regioes_geograficas.xlsx").stat().st_size, "bytes")
